from __future__ import annotations

import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from autoflow_core import detect_image_backend, evaluation_summary  # noqa: E402
from word_engine import reusable_report, validation_cache_key  # noqa: E402


def load_presentation_adapter():
    path = ROOT / "integrations" / "presentation-skill" / "scripts" / "presentation_adapter.py"
    spec = importlib.util.spec_from_file_location("autoflow_presentation_adapter", path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class EvaluationContractTests(unittest.TestCase):
    def test_eval_prompts_are_valid_utf8_chinese(self) -> None:
        payload = json.loads((ROOT / "evals" / "evals.json").read_text(encoding="utf-8"))
        self.assertEqual([item["id"] for item in payload["evals"]], [1, 2, 3, 4, 5])
        for item in payload["evals"]:
            prompt = item["prompt"]
            self.assertNotIn("\ufffd", prompt)
            self.assertTrue(any("\u4e00" <= char <= "\u9fff" for char in prompt), prompt)
            self.assertGreaterEqual(len(item.get("assertions", [])), 3)

    def test_checkpoint_pass_is_not_full_test_pass(self) -> None:
        workflow = {
            "workflow_id": "wf-eval",
            "recipe": "document",
            "steps": [
                {
                    "id": "word",
                    "module": "word",
                    "action": "fill",
                    "optional": False,
                    "outputs": ["word.document", "word.validation"],
                }
            ],
        }
        state = {
            "status": "awaiting_delivery_approval",
            "steps": {"word": {"status": "completed"}},
            "gates": {
                "plan": {"active": False, "status": "approved"},
                "source": {"active": False, "status": "not_applicable"},
                "visual": {"active": False, "status": "not_applicable"},
                "delivery": {"active": True, "status": "pending"},
            },
        }
        manifest = {"artifacts": [{"id": "word.document"}, {"id": "word.validation"}]}
        report = evaluation_summary(workflow, state, manifest, expected_gate="delivery")
        self.assertEqual(report["outcome"], "checkpoint_pass")
        self.assertTrue(report["artifact_execution_complete"])
        self.assertFalse(report["full_test_eligible"])

        state["status"] = "completed"
        state["gates"]["delivery"] = {"active": False, "status": "approved"}
        report = evaluation_summary(workflow, state, manifest)
        self.assertEqual(report["outcome"], "full_test_pass")
        self.assertTrue(report["full_test_eligible"])

    @patch("autoflow_core._image_action_report")
    def test_image_backend_reports_partial_without_sibling_leakage(self, mocked_action) -> None:
        mocked_action.side_effect = lambda action, _root: {
            "action": action,
            "status": "blocked" if action == "ai" else "available",
        }
        report = detect_image_backend(["ai", "capture", "diagram"])
        self.assertEqual(report["status"], "partial")
        self.assertEqual(report["message"], "ai=blocked, capture=available, diagram=available")

    def test_word_validation_cache_requires_exact_inputs(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            document = root / "document.docx"
            template = root / "template.docx"
            report_path = root / "report.json"
            document.write_bytes(b"document-v1")
            template.write_bytes(b"template-v1")
            key = validation_cache_key(document, template, None)
            report_path.write_text(
                json.dumps(
                    {
                        "$schema": "autoflow/word-core-validation/1.0",
                        "overall_pass": True,
                        "cache": {"input_key": key, "hit": False},
                    }
                ),
                encoding="utf-8",
            )
            self.assertIsNotNone(reusable_report(report_path, key))
            document.write_bytes(b"document-v2")
            self.assertIsNone(reusable_report(report_path, validation_cache_key(document, template, None)))

    def test_presentation_render_cache_invalidates_on_deck_change(self) -> None:
        adapter = load_presentation_adapter()
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            deck = root / "deck.pptx"
            render_dir = root / "renders"
            render_dir.mkdir()
            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[6])
            presentation.save(deck)
            (render_dir / "slide-1.png").write_bytes(b"png")
            adapter._write_render_cache(deck, render_dir, 1)
            self.assertTrue(adapter._render_cache_valid(deck, render_dir))
            deck.write_bytes(deck.read_bytes() + b"changed")
            self.assertFalse(adapter._render_cache_valid(deck, render_dir))


if __name__ == "__main__":
    unittest.main()
